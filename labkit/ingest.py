"""Ingest real panel + sample spreadsheets from an experiment's data folder, so a backfilled
experiment inherits its actual antibody/metal panel and patient list (not just placeholders)."""
from __future__ import annotations

import glob
import os

try:
    import openpyxl
except ImportError:  # optional; ingestion just no-ops without it
    openpyxl = None


def _rows(path: str) -> list[tuple]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return list(wb.active.iter_rows(values_only=True))


def _col(header: list[str], *names) -> int | None:
    for i, h in enumerate(header):
        for n in names:
            if n.lower() in (h or "").lower():
                return i
    return None


def read_panels(data_dir: str) -> dict[str, list[str]]:
    """{panel_file_stem: [antigens]} for every *panel*.xlsx (CATALYST fcs_colname/antigen layout)."""
    out = {}
    if openpyxl is None:
        return out
    for f in sorted(glob.glob(os.path.join(data_dir, "**", "*panel*.xlsx"), recursive=True)):
        try:
            rows = _rows(f)
        except Exception:
            continue
        if not rows:
            continue
        header = [str(c).strip().lower() if c else "" for c in rows[0]]
        ai = _col(header, "antigen") or (1 if len(header) > 1 else 0)
        markers = [str(r[ai]).strip() for r in rows[1:] if ai < len(r) and r[ai]]
        if markers:
            out[os.path.splitext(os.path.basename(f))[0]] = markers
    return out


def read_samples(data_dir: str) -> dict:
    """{file, samples:[{id, diagnosis, sex, age}]} from the first *sample*.xlsx found."""
    if openpyxl is None:
        return {"file": None, "samples": []}
    for f in sorted(glob.glob(os.path.join(data_dir, "**", "*sample*.xlsx"), recursive=True)):
        try:
            rows = _rows(f)
        except Exception:
            continue
        if not rows:
            continue
        header = [str(c).strip() if c else "" for c in rows[0]]
        id_i = _col(header, "sample id", "sample")
        dx_i = _col(header, "diagnosis", "diagnos")
        sex_i = _col(header, "sex")
        age_i = _col(header, "age")
        if id_i is None:
            id_i = 0
        samples = []
        for r in rows[1:]:
            sid = r[id_i] if id_i < len(r) else None
            if sid is None or str(sid).strip() == "":
                continue
            samples.append({
                "id": str(sid).strip(),
                "diagnosis": str(r[dx_i]).strip() if dx_i is not None and dx_i < len(r) and r[dx_i] else "",
                "sex": str(r[sex_i]).strip() if sex_i is not None and sex_i < len(r) and r[sex_i] else "",
                "age": str(r[age_i]).strip() if age_i is not None and age_i < len(r) and r[age_i] else "",
            })
        return {"file": os.path.basename(f), "samples": samples}
    return {"file": None, "samples": []}


def panel_block(panels: dict[str, list[str]]) -> str:
    if not panels:
        return "_Antibody/metal panel, sample list, conditions. Lives in the protocol sheet; summarise here._"
    lines = []
    for name, markers in panels.items():
        lines.append(f"- **{name}** ({len(markers)} markers): {', '.join(markers)}")
    return "Ingested from the run's panel sheets:\n\n" + "\n".join(lines)


def sample_block(s: dict) -> str:
    if not s.get("samples"):
        return ""
    lines = [f"Ingested from `{s['file']}` ({len(s['samples'])} samples):", "",
             "| Sample | Diagnosis | Sex | Age |", "|---|---|---|---|"]
    for x in s["samples"]:
        lines.append(f"| {x['id']} | {x['diagnosis']} | {x['sex']} | {x['age']} |")
    return "\n".join(lines)

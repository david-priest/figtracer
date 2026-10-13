"""figtracer.deck and figtracer.embeds: a deck built with provenance on every figure.

The experiment is synthetic (EXP, D1, treated). Renders are PNGs so no pdftoppm or Chrome
is needed; the one test that builds a real deck skips without python-pptx.
"""
from __future__ import annotations

import io
import json
import os

import pytest
from lxml import etree

from figtracer import deck as deckmod
from figtracer import embeds, manifest, ooxml
from figtracer.ooxml import q
from tests.ooxml_fixtures import _png, _pptx

OLD = "2026-01-01_EXP/2026-01-01_10.00.00_UMAP_merged.png"
NEW = "2026-01-02_EXP/2026-01-02_10.00.00_UMAP_merged.png"


def _png_margin(w: int, h: int, colour=(200, 30, 30), inset: int = 10) -> bytes:
    """A white page with a coloured drawing inset by `inset` pixels, for the trim."""
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (w, h), (255, 255, 255))
    ImageDraw.Draw(im).rectangle([inset, inset, w - inset - 1, h - inset - 1], fill=colour)
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


def _line(outputs, **fields):
    with open(os.path.join(outputs, manifest.MANIFEST_NAME), "a", encoding="utf-8") as fh:
        fh.write(json.dumps(fields) + "\n")


def _exp(tmp_path, name="EXP"):
    """An experiment with two renders of UMAP_merged (the newer one written by Python, with
    microseconds, the older by R with a zone: `saved_at_key` must order them), one render
    of paper_umap, and one table."""
    root = tmp_path / name
    outputs = root / "outputs"
    outputs.mkdir(parents=True)
    (root / ".here").write_text("")
    (root / "deck").mkdir()
    for rel, png, saved in ((OLD, _png_margin(60, 40), "2026-01-01T10:00:00+0900"),
                            (NEW, _png_margin(60, 40, (30, 200, 30)), "2026-01-02T10:00:00.123456")):
        p = outputs / rel
        p.parent.mkdir(exist_ok=True)
        p.write_bytes(png)
        _line(outputs, fig=os.path.basename(rel), rel_path=rel, title="UMAP_merged", channel="note",
              embed=True, fig_format="png", timestamp=rel.split("/")[1][:19], saved_at=saved,
              qmd_path="/x/EXP.qmd", chunk_label="fig-umap", git_commit="abc123")
    rel = "2026-01-02_EXP/2026-01-02_11.00.00_paper_umap.png"
    (outputs / rel).write_bytes(_png(10, 10, (0, 0, 200)))
    _line(outputs, fig=os.path.basename(rel), rel_path=rel, title="paper_umap", channel="note", embed=True,
          fig_format="png", timestamp="2026-01-02_11.00.00", saved_at="2026-01-02T11:00:00")
    (outputs / "population_numbers.csv").write_text("population,cells\ntreated,12\ncontrol,30\n")
    _line(outputs, kind="table", fig="population_numbers.csv", rel_path="population_numbers.csv",
          title="population_numbers", channel="note", embed=True, fig_format="csv", n_rows=2, n_cols=2,
          columns=["population", "cells"], timestamp="2026-01-02_12.00.00", saved_at="2026-01-02T12:00:00")
    return root


class _Shape:
    """Enough of a python-pptx picture to be tagged: an lxml element with nvPicPr/cNvPr."""

    def __init__(self):
        self._element = etree.fromstring(
            f'<p:pic xmlns:p="{ooxml.NS["p"]}"><p:nvPicPr><p:cNvPr id="4" name="Picture 4" descr="x.png"/>'
            f'<p:cNvPicPr/><p:nvPr/></p:nvPicPr></p:pic>')
        self.name = "Picture 4"


# ── resolving and caching figures ─────────────────────────────────────────────

def test_figure_takes_the_newest_render_by_saved_at_key_and_caches_it(tmp_path):
    root = _exp(tmp_path)
    build = deckmod.DeckBuild(str(root), title="results_for_D1", generator="deck/build.py")
    res = build.figure("UMAP_merged", trim=False)
    assert res.entry["rel_path"] == NEW                          # not the R-spelled older one
    assert res.png == str(root / "deck" / "png" / "UMAP_merged.png") and (res.width, res.height) == (60, 40)
    sidecar = json.load(open(res.png + ".json"))
    assert sidecar["rel_path"] == NEW and sidecar["dpi"] == 300 and sidecar["trim"] is False
    first = os.path.getmtime(res.png)
    build2 = deckmod.DeckBuild(str(root), title="results_for_D1")
    build2.figure("UMAP_merged", trim=False)
    assert os.path.getmtime(res.png) == first                    # same render: the cache is reused
    assert build.sources.as_list() == [{"kind": "figure", "exp": "EXP", "title": "UMAP_merged",
                                        "rel_path": NEW, "saved_at": "2026-01-02T10:00:00.123456"}]


def test_figure_rerenders_when_a_newer_entry_appears(tmp_path):
    root = _exp(tmp_path)
    build = deckmod.DeckBuild(str(root), title="results_for_D1")
    png = build.figure("UMAP_merged", trim=False).png
    newer = "2026-01-03_EXP/2026-01-03_10.00.00_UMAP_merged.png"
    (root / "outputs" / newer).parent.mkdir()
    (root / "outputs" / newer).write_bytes(_png(8, 8, (9, 9, 9)))
    _line(root / "outputs", fig=os.path.basename(newer), rel_path=newer, title="UMAP_merged",
          fig_format="png", timestamp="2026-01-03_10.00.00", saved_at="2026-01-03T10:00:00")
    res = deckmod.DeckBuild(str(root), title="results_for_D1").figure("UMAP_merged", trim=False)
    assert res.entry["rel_path"] == newer and (res.width, res.height) == (8, 8)
    assert json.load(open(png + ".json"))["rel_path"] == newer


def test_figure_trims_the_page_margin_and_keeps_a_border(tmp_path):
    root = _exp(tmp_path)
    res = deckmod.DeckBuild(str(root), title="results_for_D1").figure("UMAP_merged", border=5)
    assert (res.width, res.height) == (60 - 20 + 10, 40 - 20 + 10)


def test_figure_names_what_is_missing(tmp_path):
    root = _exp(tmp_path)
    build = deckmod.DeckBuild(str(root), title="results_for_D1")
    with pytest.raises(KeyError, match="UMAP_merged"):            # the nearest title is offered
        build.figure("UMAP_merge")
    os.remove(root / "outputs" / NEW)
    os.remove(root / "outputs" / OLD)
    with pytest.raises(FileNotFoundError):
        build.figure("UMAP_merged")
    with pytest.raises(KeyError):
        build.figure("UMAP_merged", exp="OTHER")


def test_table_rows_and_source_hash(tmp_path):
    root = _exp(tmp_path)
    build = deckmod.DeckBuild(str(root), title="results_for_D1")
    rows = build.table("population_numbers")
    assert rows == [{"population": "treated", "cells": "12"}, {"population": "control", "cells": "30"}]
    src = build.sources.as_list()[0]
    assert src["kind"] == "table" and src["sha256"] == ooxml.sha256(str(root / "outputs" / "population_numbers.csv"))
    with pytest.raises(KeyError):
        build.table("population_number")


def test_tag_picture_and_tag_table_write_the_shape_properties(tmp_path):
    root = _exp(tmp_path)
    build = deckmod.DeckBuild(str(root), title="results_for_D1")
    shape = _Shape()
    build.tag_picture(shape, "UMAP_merged")
    cnv = shape._element.find(q("p", "nvPicPr") + "/" + q("p", "cNvPr"))
    assert cnv.get("name") == "figtracer:EXP:UMAP_merged"
    assert ooxml.parse_tag(cnv.get("name"), cnv.get("descr"))["rel_path"] == NEW
    tbl = _Shape()
    build.tag_table(tbl, "population_numbers")
    cnv = tbl._element.find(q("p", "nvPicPr") + "/" + q("p", "cNvPr"))
    assert cnv.get("name") == "figtracer-table:EXP:population_numbers"
    with pytest.raises(TypeError):
        build.tag_picture(object(), "UMAP_merged")


# ── the record and the guard ──────────────────────────────────────────────────

def _built_deck(root, name="results.pptx"):
    tag = ooxml.make_tag("EXP", "UMAP_merged", {"rel_path": NEW, "saved_at": "2026-01-02T10:00:00.123456"})
    return _pptx(root / "deck", [{"texts": [("Title", "D1")], "pics": [(tag[0], tag[1], "m")]}],
                 name=name, media={"m": _png(4, 4)})


def test_record_appends_the_line_sets_properties_and_the_gitignore_once(tmp_path):
    root = _exp(tmp_path)
    (root / ".gitignore").write_text("analysis/*.qs2\n")
    build = deckmod.DeckBuild(str(root), title="results_for_D1", generator="deck/build.py --full")
    build.figure("UMAP_merged")
    build.table("population_numbers")
    path = _built_deck(root)
    rec = build.record(path)
    assert rec["kind"] == "deck" and rec["rel_path"] == "../deck/results.pptx" and rec["fig_format"] == "pptx"
    assert rec["n_units"] == 1 and rec["generator"] == "deck/build.py --full" and rec["tool"] == "figtracer.deck"
    assert [s["title"] for s in rec["sources"]] == ["UMAP_merged", "population_numbers"]
    assert rec["sha256"] == ooxml.sha256(path)                   # the hash of the file as it leaves
    props = ooxml.core_properties(path)
    assert props["subject"] == "EXP" and props["identifier"] == "results_for_D1"
    assert props["description"].startswith("Built ") and "1 figure(s), 1 table(s)" in props["description"]
    lines = manifest.read_entries(str(root / "outputs" / manifest.MANIFEST_NAME))
    assert manifest.kind(lines[-1]) == "deck" and lines[-1]["title"] == "results_for_D1"
    gi = (root / ".gitignore").read_text()
    assert gi.startswith("analysis/*.qs2\n") and "deck/*.pptx\n" in gi and "deck/png/\n" in gi
    build.record(path)
    assert (root / ".gitignore").read_text().count("deck/*.pptx") == 1
    assert embeds.latest(str(root / "outputs"), "results_for_D1")["sha256"] == ooxml.sha256(path)
    assert embeds.record_of_file(str(root / "outputs"), path)["title"] == "results_for_D1"


def test_target_diverts_when_the_deck_is_open_or_edited(tmp_path):
    root = _exp(tmp_path)
    outputs = str(root / "outputs")
    build = deckmod.DeckBuild(str(root), title="results_for_D1")
    assert build.target("results.pptx") == str(root / "deck" / "results.pptx")
    path = _built_deck(root)
    assert build.target("results.pptx") == path                  # exists, never recorded: writable
    build.record(path)
    assert build.target("results.pptx") == path                  # recorded and unchanged
    ooxml.set_core_properties(path, description="a hand edit")
    assert build.target("results.pptx") == str(root / "deck" / "results v2.pptx")
    assert embeds.target(path, outputs, "results_for_D1")[1].startswith("results.pptx is edited since")
    (root / "deck" / "~$results.pptx").write_bytes(b"")
    assert embeds.target(path, outputs, "results_for_D1")[1].startswith("results.pptx is open in PowerPoint")


# ── status ────────────────────────────────────────────────────────────────────

def _status_deck(root):
    stale = ooxml.make_tag("EXP", "UMAP_merged", {"rel_path": OLD, "saved_at": "2026-01-01T10:00:00+0900"})
    ok = ooxml.make_tag("EXP", "paper_umap", {"rel_path": "2026-01-02_EXP/2026-01-02_11.00.00_paper_umap.png"})
    gone = ooxml.make_tag("EXP", "no_such_figure", {"rel_path": "x.png"})
    other = ooxml.make_tag("OTHER", "paper_umap", {"rel_path": "y.png"})
    tbl = ooxml.make_tag("EXP", "population_numbers", {"sha256": "not the hash"}, table=True)
    return _pptx(root / "deck", [
        {"texts": [("Title", "D1")], "pics": [(stale[0], stale[1], "a"), ("Picture 9", "hand.png", "b")]},
        {"texts": [("Title", "two")], "pics": [(ok[0], ok[1], "a"), (gone[0], gone[1], "b"), (other[0], other[1], "a")],
         "table": (tbl[0], tbl[1], [["population", "cells"], ["treated", "12"]])},
    ], name="status.pptx", media={"a": _png(4, 4), "b": _png(6, 6)})


def test_embed_status_and_file_status(tmp_path):
    root = _exp(tmp_path)
    outputs = str(root / "outputs")
    path = _status_deck(root)
    resolver = embeds.scope_resolver("EXP", str(root), True)
    rows = embeds.embed_status(ooxml.read(path), resolver)
    states = [(r["kind"], r["where"], r["state"]) for r in rows]
    assert states == [("figure", "slide 1", "STALE"), ("figure", "slide 1", "UNTRACKED"),
                      ("figure", "slide 2", "ok"), ("figure", "slide 2", "DANGLING"),
                      ("figure", "slide 2", "UNKNOWN EXP"), ("table", "slide 2", "STALE")]
    assert rows[0]["detail"] == "built from 2026-01-01T10:00, newest 2026-01-02T10:00"
    rec = embeds.record(path, outputs=outputs, title="status", scope="EXP")
    e = embeds.resolve_documents(outputs)["status"]
    assert embeds.file_status(e) == "ok"
    ooxml.set_core_properties(path, description="edited by hand")
    assert embeds.file_status(e) == "EDITED SINCE RECORD"
    (root / "deck" / "~$status.pptx").write_bytes(b"")
    assert embeds.file_status(e) == "OPEN IN POWERPOINT"
    os.remove(path)
    assert embeds.file_status(e) == "MISSING"
    assert rec["sources"][0]["title"] == "UMAP_merged"           # adopted from the tags


def test_refresh_picture_installs_the_newest_render(tmp_path):
    root = _exp(tmp_path)
    path = _status_deck(root)
    doc = ooxml.read(path)
    figs, _ = embeds.scope_resolver("EXP", str(root), True)("EXP")
    report = embeds.refresh_picture(path, doc.pictures[0], figs["UMAP_merged"])
    assert report["ok"] and report["skipped"] is None
    after = ooxml.read(path)
    assert after.pictures[0].tag["rel_path"] == NEW
    assert ooxml.read_part(path, after.pictures[0].media) == _png_margin(60, 40, (30, 200, 30))
    assert after.pictures[2].media != after.pictures[0].media  # the part slide 2 shared is untouched
    rows = embeds.embed_status(after, embeds.scope_resolver("EXP", str(root), True))
    assert rows[0]["state"] == "ok"


# ── the command ───────────────────────────────────────────────────────────────

def test_cli_check_record_and_inspect(tmp_path, monkeypatch, capsys):
    root = _exp(tmp_path)
    outputs = str(root / "outputs")
    monkeypatch.setattr(embeds, "scope", lambda args: ("EXP", outputs, str(root), True))
    path = _status_deck(root)
    assert deckmod.main(["record", path, "--exp", "EXP", "--title", "status", "--generator", "by hand"]) == 0
    out = capsys.readouterr().out
    assert "recorded status" in out and "4 figure(s), 1 table(s)" in out
    assert deckmod.main(["check", "--exp", "EXP"]) == 1
    out = capsys.readouterr().out
    assert "== status" in out and "[STALE FIGURE]  slide 1  EXP:UMAP_merged" in out
    assert "[UNTRACKED FIGURE]  slide 1  hand.png" in out and "[DANGLING FIGURE]  slide 2" in out
    assert "[UNKNOWN EXP FIGURE]  slide 2  OTHER:paper_umap" in out and "[STALE TABLE]  slide 2" in out
    assert "[ok]  slide 2  EXP:paper_umap" in out
    assert deckmod.main(["inspect", path, "--exp", "EXP"]) == 0
    out = capsys.readouterr().out
    assert "status.pptx: pptx, 2 slides, 5 picture(s), 1 table(s)" in out
    assert "record          " in out and "as status, the file on disk, by by hand" in out
    assert "slide 1    figure  EXP:UMAP_merged" in out and "picture 2: untracked (hand.png)" in out
    assert deckmod.main(["check", "--exp", "EXP", "--title", "nothing"]) == 2
    assert deckmod.main(["record", path, "--title", "x"]) == 2    # no scope


# ── the real thing ────────────────────────────────────────────────────────────

def test_a_deck_built_with_python_pptx_reads_back_tagged(tmp_path, capsys):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches
    root = _exp(tmp_path)
    build = deckmod.DeckBuild(str(root), title="results_for_D1", generator="deck/build.py", dpi=150)
    prs = pptx.Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    pic = build.add_figure(s, "UMAP_merged", Inches(0.4), Inches(0.8), Inches(8), Inches(2))
    assert pic.name == "figtracer:EXP:UMAP_merged" and pic.height <= Inches(2)
    build.add_figure(s, "paper_umap", Inches(9), Inches(0.8), Inches(3))
    rows = build.table("population_numbers")
    shape = s.shapes.add_table(len(rows) + 1, 2, Inches(0.5), Inches(4), Inches(6), Inches(2))
    shape.table.cell(0, 0).text, shape.table.cell(0, 1).text = "population", "cells"
    for i, r in enumerate(rows, 1):
        shape.table.cell(i, 0).text, shape.table.cell(i, 1).text = r["population"], r["cells"]
    build.tag_table(shape, "population_numbers")
    out = build.target("results.pptx")
    prs.save(out)
    rec = build.record(out)
    doc = ooxml.read(out)
    assert doc.n_units == 1 and [p.tag["title"] for p in doc.pictures] == ["UMAP_merged", "paper_umap"]
    assert doc.pictures[0].tag["rel_path"] == NEW and doc.pictures[0].tag["dpi"] == 150
    assert doc.tables[0].tag["title"] == "population_numbers" and doc.tables[0].rows[1] == ["treated", "12"]
    assert "12" in [p.text for p in doc.paragraphs]
    rows = embeds.embed_status(doc, embeds.scope_resolver("EXP", str(root), True))
    assert [r["state"] for r in rows] == ["ok", "ok", "ok"]
    assert rec["sha256"] == ooxml.sha256(out) and ooxml.core_properties(out)["identifier"] == "results_for_D1"
    assert deckmod.main(["inspect", out]) == 0
    assert "2 picture(s), 1 table(s)" in capsys.readouterr().out
